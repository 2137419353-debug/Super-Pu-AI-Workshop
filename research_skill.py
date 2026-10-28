# -*- coding: utf-8 -*-
"""研究黑盒（从 skill-searcher_v02 提取，仅保留无 UI 的核心链路）。

``run_skill(topic, api_key)`` 是唯一后端入口：校验输入 → 组装 ``SKILL_research.md``
Prompt → 调 DeepSeek → 清洗 / 八维结构校验 / 一次格式修复 → Markdown 转 Word。
返回 ``{"text": str, "mermaid_code": str, "word_bytes": bytes}``，不含任何 Streamlit 代码。

与原始单机版的差异只有三处（均为接入协作空间所必需）：
1. 去掉 Streamlit / components 导入与全部 UI 渲染代码；
2. Prompt 文件名由 ``SKILL.md`` 改为 ``SKILL_research.md``（避免与 Pre 技能的根 SKILL.md 重名）；
3. 末尾补充 ``report_preview`` / ``report_filename`` / ``report_stats`` 三个只读辅助函数，
   供网页层渲染左侧卡片、写回成果篮时复用，不改变 ``run_skill`` 行为。
"""

from __future__ import annotations

import io
import json
import os
import re
from datetime import date
from pathlib import Path
from typing import Any, Final, TypedDict
from urllib import error as urlerror
from urllib import request as urlrequest

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from docx.text.run import Run

# ---------------------------------------------------------------------------
# 常量与类型
# ---------------------------------------------------------------------------

APP_TITLE: Final = "关键词深度调研舱"
APP_ICON: Final = "🔎"
PROJECT_ROOT: Final = Path(__file__).resolve().parent
SKILL_PATH: Final = PROJECT_ROOT / "SKILL_research.md"

DEFAULT_API_URL: Final = "https://api.deepseek.com/chat/completions"
DEFAULT_MODEL: Final = "deepseek-chat"
MAX_TOKENS: Final = 8192
MAX_TOPIC_CHARS: Final = 12000
MAX_FORMAT_RETRIES: Final = 1
DOCX_MIME: Final = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

BODY_EAST_ASIA: Final = "宋体"
HEADING_EAST_ASIA: Final = "微软雅黑"
BODY_ASCII: Final = "Times New Roman"
HEADER_FILL: Final = "D9E2F3"
BAND_FILL: Final = "F7F9FC"


class SkillError(RuntimeError):
    """可以直接展示给用户的业务错误。"""


class SkillResult(TypedDict):
    text: str
    mermaid_code: str
    word_bytes: bytes


def _env_int(name: str, default: int) -> int:
    try:
        return int(os.getenv(name, str(default)))
    except (TypeError, ValueError):
        return default


API_URL: Final = (
    os.getenv("DEEPSEEK_API_URL", DEFAULT_API_URL).strip() or DEFAULT_API_URL
)
MODEL: Final = os.getenv("DEEPSEEK_MODEL", DEFAULT_MODEL).strip() or DEFAULT_MODEL
API_TIMEOUT_SECONDS: Final = _env_int("DEEPSEEK_TIMEOUT", 300)


# ---------------------------------------------------------------------------
# Prompt：完整 SKILL.md + 强制输出契约
# ---------------------------------------------------------------------------

OUTPUT_CONTRACT: Final = r"""
---

# 最高优先级：强制输出契约

【报告生成日期】{{TODAY}}

本契约覆盖 `SKILL.md` 中任何与“追问、澄清、分步交互、JSON 文件、脚本运行、工具调用”有关的操作指令。`SKILL.md` 保留完整的研究方法论、六个来源类别与写作标准；你只负责一次性输出最终的 Markdown 报告正文。

## 0. 绝对红线

1. 不得问候、寒暄、解释、道歉、自我评价、描述过程、输出免责声明或结尾祝福。
2. 不得说“好的”“以下是”“根据您的要求”“希望这份报告对您有帮助”“如需调整请告诉我”。
3. 不得向用户提出任何问题，不得要求补充信息，不得分步输出，不得等待确认。
4. 不得输出“作为 AI”“我无法……”“请注意……”之类的元话语。无法确认的信息，用可信度标签与“待核验”处理。
5. 不得把整篇报告包在 ```markdown 代码围栏中；全文只允许出现一个 ```mermaid 代码围栏。
6. 不得输出文件路径、脚本命令、JSON、present_files、WebSearch / WebFetch 调用记录。
7. 输出第一个字符必须是 `#`，最后一个字符必须是来源条目的 URL 或“URL 待核验”。
8. 如果你不按照规定格式输出，或者出现追问、废话，系统将崩溃。必须严格输出。

## 1. 唯一允许的文档骨架

严格使用下列标题文字与层级。除骨架中的标题外，不得增加新的 `#` 或 `##` 顶层标题；`####` 子标题可以按内容增加。

# [主标题]
## [副标题]
关键词：[关键词1], [关键词2], [关键词3]

> **摘要**：[150-250 字，直接给结论，不写“本文将……”]

### 一、来源出处
#### 1. [子标题]
[概念的最早出处、词源/翻译史、关键人物与原始文献；2-4 段]

### 二、含义变迁
[先用 1-2 段解释含义如何变化，再给时间轴表格]
| 时间 | 阶段 | 含义/用法变化 | 关键证据与来源 | 可信度 |
|---|---|---|---|---|
| ... | ... | ... | ... | [可信度：高 ✅] |

### 三、相关新闻与社会事件
#### 3.1 [案例一]
[真实案例：时间、地点/主体、经过、结果、与主题的关系；段末给可信度]
#### 3.2 [案例二]
...
#### 3.3 [案例三]
...

### 四、统计数据
#### 4.1 [数据表一标题]
[说明口径、年份、样本范围]
| 指标 | 数值 | 年份 | 统计口径/来源 | 可信度 |
|---|---|---|---|---|
| ... | ... | ... | ... | [可信度：中 ⚠️] |

#### 4.2 [数据表二标题]
[第二张表必须与第一张表的口径或维度不同]
| 指标 | 数值 | 年份 | 统计口径/来源 | 可信度 |
|---|---|---|---|---|

[趋势解读：至少 2 段，指出上升/下降/分化及其可能原因，不得只复述数字]

### 五、学术文献与行业报告
#### 5.1 [文献/报告一]
[作者、年份、核心问题、方法、结论、争议点；给出出处]
#### 5.2 [文献/报告二]
...
[至少 4 项，学术论文与行业报告都要覆盖]

### 六、哲学思想关联
#### 6.1 [思想传统一]
[概念如何与该传统发生关系；至少 2 个传统，不做人名堆砌]
#### 6.2 [思想传统二]
...

### 七、关键词关联分析
#### 7.1 关系网络
[用 1-2 段解释 5 个以上关联词，说明“因果/对立/上位/下位/共现”关系]
#### 7.2 关键词关系图
[先用一句话说明图的结构，再输出唯一一个 Mermaid 图；不要写“下图”，写“关系图见网页端可视化”]
```mermaid
flowchart LR
    A["核心关键词"] --> B["关联维度"]
```
### 八、辩论要点速查
#### 8.1 正反方速查
| 立场 | 核心主张 | 关键论据 | 可信度 |
|---|---|---|---|
| 正方 | ... | ... | [可信度：高 ✅] |
| 反方 | ... | ... | [可信度：中 ⚠️] |

#### 8.2 核心交锋点
| 交锋点 | 正方 | 反方 | 裁判/裁决要点 |
|---|---|---|---|

#### 8.3 金句弹药库
| 金句 | 适用场景 | 出处/说明 |
|---|---|---|

### 数据来源
[1] 来源名. 文章名. URL.
[2] 来源名. 报告名. 年份. URL.
[至少 6 条，按正文出现顺序编号]

## 2. 硬性数量与质量要求

- 全文 4500-6000 个汉字（含表格文字），不得用空话凑字数。
- 至少 3 个真实案例；每个案例必须有具体时间、主体、经过、结果、与主题的关联。
- 第四部分至少 2 张统计表，每张至少 3 行数据；必须写出趋势解读。
- 第八部分必须有 3 张表：正反方表、交锋点表、金句表。
- 第五部分至少 4 项文献/报告；第六部分至少 2 个哲学传统；第七部分至少 5 个关联词。
- 来源列表至少 6 条；正文中出现的每个关键数字都必须在来源列表中有对应条目。
- 所有表格必须是标准 Markdown 管道表：表头、分隔行、数据行齐全；单元格内不得出现未转义的 `|`。

## 3. 事实核查标签（强制）

在每一个案例段落末尾、每一行统计数据后、每一组关键事实后，紧跟以下三种标签之一：

- `[可信度：高 ✅]`：政府/国际组织统计、同行评审论文、可交叉验证的原始文件。
- `[可信度：中 ⚠️]`：主流媒体、单一机构报告、样本有限但方法透明的调查。
- `[可信度：低 ❌]`：自媒体、轶事、无法核验的网传说法、模型记忆不确定的内容。

规则：

1. 不得为了显得权威而把不确定信息标成“高”。
2. 不得编造精确到小数点的数字、日期、论文标题、作者或 URL。
3. 无法确认链接时写“URL 待核验”，不得编造假 URL。
4. 估算值必须写“估算”，并说明估算方法。
5. 同一段内事实与观点要分开；观点用“支持者认为/批评者指出”，不得伪装成事实。

## 4. 写作风格

- 简体中文，学术但不晦涩；句子短，信息密度高。
- 每个 `####` 小节先给结论，再给证据，再给限定条件。
- 不使用“首先/其次/最后”式流水账；用真实因果与争议组织内容。
- 不写“综上所述”“让我们”“可以看到”等填充语。
- 引用原文时使用引号并注明出处；转述时不要伪造直接引语。
- 对争议话题保持双边对称：正反方论据强度相当，不替用户下价值判断。

## 5. 输出前自检（不要输出自检过程）

- 是否只有一个 `#`、一个 `##`，且标题文字与骨架一致？
- 八个 `###` 与 `### 数据来源` 是否齐全？
- 第二、四、八部分表格数量是否达标？
- 第七部分是否有且只有一个 Mermaid 图？
- 案例和统计是否都有可信度标签？
- 是否存在任何问候、追问、结尾废话、代码围栏或 JSON？

现在开始直接输出终稿。不要解释，不要确认，不要追问。
"""


def _read_skill_markdown() -> str:
    try:
        text = SKILL_PATH.read_text(encoding="utf-8")
    except OSError as exc:  # pragma: no cover - 部署缺文件时才会触发
        raise SkillError(
            f"缺少 Prompt 文件：{SKILL_PATH.name}。请确认它和 app.py 在同一目录。"
        ) from exc
    if len(text) < 1000:
        raise SkillError("SKILL_research.md 内容异常（过短），请重新复制原始文件。")
    return text


def build_system_prompt() -> str:
    """完整拼装 System Prompt；SKILL_research.md 原文不做任何裁剪。"""
    skill_markdown = _read_skill_markdown()
    contract = OUTPUT_CONTRACT.replace("{{TODAY}}", date.today().isoformat())
    return f"{skill_markdown}\n\n{contract}"


# ---------------------------------------------------------------------------
# 输入校验与 API 调用
# ---------------------------------------------------------------------------


def validate_topic(user_input: str) -> str | None:
    """仅拦截空输入与极端超长输入；短关键词直接放行。"""
    compact = re.sub(r"\s+", "", user_input or "")

    if not compact:
        return "请先输入一个关键词或研究线索。"
    if len(compact) > MAX_TOPIC_CHARS:
        return (
            f"输入过长（{len(compact)} 字）。请压缩到 {MAX_TOPIC_CHARS} 字以内，"
            "保留关键词和最关键的研究角度。"
        )
    return None


def _redact(value: str, api_key: str) -> str:
    value = (value or "").replace(api_key, "sk-***") if api_key else (value or "")
    return re.sub(r"sk-[A-Za-z0-9_-]{8,}", "sk-***", value)


def _http_error_body(exc: urlerror.HTTPError) -> str:
    try:
        raw = exc.read(4096)
        text = raw.decode("utf-8", errors="replace")
    except Exception:  # pragma: no cover - 极端网络错误
        return ""
    try:
        payload = json.loads(text)
        message = payload.get("error", {}).get("message")
        if isinstance(message, str) and message.strip():
            return message.strip()
    except (ValueError, AttributeError):
        pass
    return re.sub(r"<[^>]+>", " ", text).strip()


def _friendly_http_error(status: int, detail: str, api_key: str) -> str:
    detail = _redact(detail, api_key).strip()
    if status in (401, 403):
        base = "API Key 无效或没有权限。请检查侧边栏中的 DeepSeek Key。"
    elif status == 402:
        base = "DeepSeek 账户余额不足，请充值后重试。"
    elif status == 429:
        base = "请求过于频繁或并发超限，请等待 1 分钟后重试。"
    elif status == 400:
        base = "请求参数被 DeepSeek 拒绝，请缩短输入后重试。"
    elif status >= 500:
        base = "DeepSeek 服务暂时不可用，请稍后重试。"
    else:
        base = f"DeepSeek API 返回错误（HTTP {status}）。"
    if detail:
        base += f" 服务端信息：{detail[:240]}"
    return base


def _post_json(payload: dict[str, Any], api_key: str, timeout: int) -> dict[str, Any]:
    body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
    req = urlrequest.Request(
        API_URL,
        data=body,
        method="POST",
        headers={
            "Authorization": f"Bearer {api_key}",
            "Content-Type": "application/json; charset=utf-8",
            "Accept": "application/json",
            "User-Agent": "skill-searcher/2.0",
        },
    )

    try:
        with urlrequest.urlopen(req, timeout=timeout) as response:
            raw = response.read()
    except urlerror.HTTPError as exc:
        detail = _http_error_body(exc)
        raise SkillError(_friendly_http_error(exc.code, detail, api_key)) from exc
    except urlerror.URLError as exc:
        reason = getattr(exc, "reason", exc)
        if isinstance(reason, TimeoutError) or "timed out" in str(reason).lower():
            raise SkillError(
                "DeepSeek 响应超时。报告较长时可能需要 2-5 分钟，请重试。"
            ) from exc
        raise SkillError(
            "无法连接 DeepSeek API。请检查网络、代理或 DEEPSEEK_API_URL 配置。"
        ) from exc
    except TimeoutError as exc:
        raise SkillError("DeepSeek 响应超时，请稍后重试。") from exc
    except OSError as exc:
        raise SkillError(f"网络请求失败：{_redact(str(exc), api_key)}") from exc

    try:
        data = json.loads(raw.decode("utf-8"))
    except (UnicodeDecodeError, ValueError) as exc:
        raise SkillError("DeepSeek 返回了无法解析的响应，请重试。") from exc
    if not isinstance(data, dict):
        raise SkillError("DeepSeek 返回格式异常，请重试。")
    return data


def _extract_content(data: dict[str, Any]) -> str:
    choices = data.get("choices")
    if not isinstance(choices, list) or not choices:
        raise SkillError("DeepSeek 没有返回候选内容，请重试。")

    first = choices[0] if isinstance(choices[0], dict) else {}
    message = first.get("message") if isinstance(first, dict) else None
    content = message.get("content") if isinstance(message, dict) else None

    if isinstance(content, list):
        parts: list[str] = []
        for item in content:
            if isinstance(item, dict) and isinstance(item.get("text"), str):
                parts.append(item["text"])
        content = "".join(parts)

    if not isinstance(content, str) or not content.strip():
        raise SkillError("DeepSeek 返回了空内容，请重试。")

    if first.get("finish_reason") == "length":
        raise SkillError(
            "模型输出达到长度上限，报告可能被截断。请缩短输入后重试。"
        )
    return content


def _request_chat(
    messages: list[dict[str, str]],
    api_key: str,
    *,
    temperature: float,
) -> str:
    payload = {
        "model": MODEL,
        "messages": messages,
        "temperature": temperature,
        "max_tokens": MAX_TOKENS,
        "stream": False,
    }
    return _extract_content(_post_json(payload, api_key, API_TIMEOUT_SECONDS))


def _request_report(topic: str, api_key: str) -> str:
    system_prompt = build_system_prompt()
    user_prompt = (
        "【研究主题 / 初稿】\n"
        f"{topic}\n\n"
        "请立即输出完整的八维 Markdown 调研报告。不要追问，不要解释。"
    )
    return _request_chat(
        [
            {"role": "system", "content": system_prompt},
            {"role": "user", "content": user_prompt},
        ],
        api_key,
        temperature=0.25,
    )


def _request_repair(
    topic: str, api_key: str, previous_output: str, issues: list[str]
) -> str:
    issue_text = "；".join(issues[:8])
    repair_prompt = (
        "【自动格式校验未通过】\n"
        f"缺失项：{issue_text}\n\n"
        "请重新输出一份完整的终稿，而不是只补缺失部分。"
        "严格使用规定的八个 `###`、`### 数据来源`、表格和唯一一个 Mermaid 图。"
        "不要解释，不要道歉，不要追问。"
    )
    return _request_chat(
        [
            {"role": "system", "content": build_system_prompt()},
            {"role": "user", "content": f"【研究主题 / 初稿】\n{topic}"},
            {"role": "assistant", "content": previous_output[:18000]},
            {"role": "user", "content": repair_prompt},
        ],
        api_key,
        temperature=0.1,
    )


# ---------------------------------------------------------------------------
# Markdown 清洗、Mermaid 提取、结构自检
# ---------------------------------------------------------------------------

_MERMAID_RE = re.compile(
    r"```(?:mermaid|mmd)\s*\n(?P<code>.*?)```", re.IGNORECASE | re.DOTALL
)
_UNLABELED_DIAGRAM_RE = re.compile(
    r"```\s*(?P<code>(?:flowchart|graph)\s+(?:TD|TB|LR|RL|BT)[\s\S]*?)```",
    re.IGNORECASE,
)
_INLINE_RE = re.compile(
    r"(?P<bold>\*\*(?P<bold_text>.+?)\*\*)"
    r"|(?P<italic>(?<!\*)\*(?!\*)(?P<italic_text>[^*\n]+?)\*(?!\*))"
    r"|(?P<code>`(?P<code_text>[^`\n]+)`)"
    r"|(?P<link>\[(?P<link_text>[^\]]+)\]\((?P<link_url>(?:https?://|mailto:)[^)\s]+)\))"
    r"|(?P<cred>\[可信度\s*[:：]\s*(?P<cred_level>高|中|低)\s*(?P<cred_icon>✅|⚠️|❌)?\])"
)
_KEYWORDS_RE = re.compile(r"^\s*(?:\*\*)?关键词(?:\*\*)?\s*[:：]\s*(?P<value>.+?)\s*$")
_ABSTRACT_RE = re.compile(
    r"(?m)^(?:>\s*)?\*\*摘要\s*[:：]?\*\*\s*[:：]?"
)
_TABLE_SEPARATOR_RE = re.compile(r"^:?-{3,}:?$")
_FLUFF_LINE_RE = re.compile(
    r"^\s*(?:如果(?:你)?需要|如需|需要我|要不要我|希望(?:这份|本|以上)|"
    r"以上(?:就是|报告|内容)|请问|是否|我可以|请告诉我|欢迎|告诉我|祝(?:你|您)|"
    r"注\s*[:：]|备注\s*[:：]|（注|\(注|全文完|报告完|——\s*完\s*——)"
)


def extract_mermaid(text: str) -> tuple[str, str]:
    """返回 (去掉图代码的正文, Mermaid 源码)。只取第一个图，移除所有图块。"""
    source = text or ""
    blocks = list(_MERMAID_RE.finditer(source))
    if not blocks:
        blocks = list(_UNLABELED_DIAGRAM_RE.finditer(source))

    if not blocks:
        return source, ""

    mermaid = blocks[0].group("code").strip()
    mermaid = re.sub(r"^\s*(?:mermaid|mmd)\s*\n", "", mermaid, flags=re.I)
    mermaid = re.sub(r"^\s*```.*$", "", mermaid, flags=re.M).strip()
    source = _MERMAID_RE.sub("", source)
    source = _UNLABELED_DIAGRAM_RE.sub("", source)
    return source, mermaid[:20000]


def _is_reference_line(line: str) -> bool:
    stripped = line.strip()
    if not stripped:
        return False
    if re.match(r"^\s*(?:[-*]\s*)?(?:\[\d+\]|\d+[.、])\s*\S", line):
        return True
    return "http://" in stripped or "https://" in stripped


def _truncate_after_data_sources(text: str) -> str:
    """保留“数据来源”标题及编号引用，砍掉其后的客套和追问。"""
    match = re.search(r"(?m)^###\s*数据来源\s*$", text)
    if not match:
        return text

    head = text[: match.end()]
    tail = text[match.end() :]
    kept: list[str] = []
    refs = 0

    for line in tail.splitlines():
        stripped = line.strip()
        if not stripped:
            kept.append("")
            continue
        if _is_reference_line(line):
            refs += 1
            kept.append(line.rstrip())
            continue
        if refs and (line.startswith((" ", "\t")) or stripped.startswith(("URL", "http"))):
            kept.append(line.rstrip())
            continue
        break

    if refs == 0:
        return head.rstrip() + "\n"
    return head.rstrip() + "\n" + "\n".join(kept).strip() + "\n"


def _strip_trailing_chatter(text: str) -> str:
    lines = text.splitlines()
    while lines:
        last = lines[-1].strip()
        if not last:
            lines.pop()
            continue
        if _FLUFF_LINE_RE.match(last) or re.fullmatch(r"[-*_—]{3,}", last):
            lines.pop()
            continue
        break
    return "\n".join(lines)


def clean_model_text(raw: str) -> str:
    """物理清洗：砍掉前言、结尾废话、整篇代码围栏与残留围栏。"""
    if not isinstance(raw, str) or not raw.strip():
        return ""

    text = raw.replace("\r\n", "\n").replace("\r", "\n").replace("\ufeff", "").strip()
    text = re.sub(r"^```(?:markdown|md|text)\s*\n", "", text, flags=re.I)
    text = re.sub(r"\n```\s*$", "", text)

    title_match = re.search(r"(?m)^#\s+\S.*$", text)
    if title_match:
        text = text[title_match.start() :]

    text = _truncate_after_data_sources(text)
    text = _strip_trailing_chatter(text)
    text = re.sub(r"(?m)^\s*[-*_]{3,}\s*$", "", text)
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    text = text.strip()
    return f"{text}\n" if text else ""


def postprocess_report(raw: str) -> tuple[str, str]:
    # 先保留代码围栏，才能识别 Mermaid；提取完再清掉剩余围栏。
    text = clean_model_text(raw)
    text, mermaid = extract_mermaid(text)
    text = re.sub(r"(?m)^\s*```.*$", "", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return (f"{text}\n" if text else "", mermaid)


def _split_table_row(row: str) -> list[str]:
    """解析管道行；支持 `\\|` 转义、行内代码中的管道与无首尾管道。"""
    text = row.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|") and not text.endswith("\\|"):
        text = text[:-1]

    cells: list[str] = []
    buffer: list[str] = []
    escaped = False
    in_code = False

    for char in text:
        if escaped:
            buffer.append(char)
            escaped = False
        elif char == "\\":
            escaped = True
        elif char == "`":
            in_code = not in_code
            buffer.append(char)
        elif char == "|" and not in_code:
            cells.append("".join(buffer).strip())
            buffer = []
        else:
            buffer.append(char)

    if escaped:
        buffer.append("\\")
    cells.append("".join(buffer).strip())
    return cells


def _is_table_separator(line: str) -> bool:
    if "|" not in line and not line.strip().startswith(":"):
        return False
    cells = _split_table_row(line)
    meaningful = [cell for cell in cells if cell.strip()]
    if not meaningful:
        return False
    return all(
        _TABLE_SEPARATOR_RE.fullmatch(cell.replace(" ", "")) for cell in meaningful
    )


def _looks_like_table_row(line: str) -> bool:
    return "|" in line and bool(line.strip())


def _iter_table_blocks(text: str) -> list[str]:
    lines = text.splitlines()
    blocks: list[str] = []
    index = 0
    while index < len(lines) - 1:
        if _looks_like_table_row(lines[index]) and _is_table_separator(lines[index + 1]):
            start = index
            index += 2
            while index < len(lines) and _looks_like_table_row(lines[index]):
                index += 1
            blocks.append("\n".join(lines[start:index]))
        else:
            index += 1
    return blocks


def _count_tables(text: str) -> int:
    return len(_iter_table_blocks(text))


def _section_text(text: str, heading_pattern: str) -> str:
    match = re.search(heading_pattern, text, flags=re.MULTILINE)
    if not match:
        return ""
    start = match.end()
    next_heading = re.search(r"(?m)^###\s+", text[start:])
    end = start + next_heading.start() if next_heading else len(text)
    return text[start:end]


SECTION_PATTERNS: Final[tuple[tuple[str, str], ...]] = (
    ("一、来源出处", r"^###\s*一\s*[、.．]?\s*来源出处\s*$"),
    ("二、含义变迁", r"^###\s*二\s*[、.．]?\s*含义变迁\s*$"),
    ("三、相关新闻与社会事件", r"^###\s*三\s*[、.．]?\s*相关新闻与社会事件\s*$"),
    ("四、统计数据", r"^###\s*四\s*[、.．]?\s*统计数据\s*$"),
    ("五、学术文献与行业报告", r"^###\s*五\s*[、.．]?\s*学术文献与行业报告\s*$"),
    ("六、哲学思想关联", r"^###\s*六\s*[、.．]?\s*哲学思想关联\s*$"),
    ("七、关键词关联分析", r"^###\s*七\s*[、.．]?\s*关键词关联分析\s*$"),
    ("八、辩论要点速查", r"^###\s*八\s*[、.．]?\s*辩论要点速查\s*$"),
    ("数据来源", r"^###\s*数据来源\s*$"),
)


def validate_report_structure(text: str, mermaid_code: str) -> list[str]:
    """返回缺失项列表；空列表表示通过。"""
    issues: list[str] = []
    h1_count = len(re.findall(r"(?m)^#\s+\S", text))
    h2_count = len(re.findall(r"(?m)^##\s+\S", text))
    if h1_count == 0:
        issues.append("缺少 `#` 主标题")
    elif h1_count > 1:
        issues.append("出现多个 `#` 一级标题")
    if h2_count == 0:
        issues.append("缺少 `##` 副标题")
    elif h2_count > 1:
        issues.append("出现多个 `##` 二级标题")
    if not re.search(r"(?m)^关键词\s*[:：]\s*\S", text):
        issues.append("缺少“关键词：”行")
    if not _ABSTRACT_RE.search(text):
        issues.append("缺少 `**摘要**`")

    for label, pattern in SECTION_PATTERNS:
        if not re.search(pattern, text, flags=re.MULTILINE):
            issues.append(f"缺少章节：{label}")

    if _count_tables(text) < 6:
        issues.append("表格总数不足（至少 6 张）")

    section_two = _section_text(text, SECTION_PATTERNS[1][1])
    if _count_tables(section_two) < 1:
        issues.append("“含义变迁”缺少时间轴表格")

    section_four = _section_text(text, SECTION_PATTERNS[3][1])
    if _count_tables(section_four) < 2:
        issues.append("“统计数据”少于 2 张数据表")

    section_eight = _section_text(text, SECTION_PATTERNS[7][1])
    if _count_tables(section_eight) < 3:
        issues.append("“辩论要点速查”少于 3 张表")

    if len(mermaid_code.strip()) < 20:
        issues.append("缺少 Mermaid 关键词关系图")

    source_section = _section_text(text, SECTION_PATTERNS[8][1])
    references = re.findall(
        r"(?m)^\s*(?:[-*]\s*)?(?:\[\d+\]|\d+[.、])\s*\S", source_section
    )
    if len(references) < 3:
        issues.append("数据来源少于 3 条编号引用")

    credibility_count = len(
        re.findall(r"\[可信度\s*[:：]\s*(?:高|中|低)", text)
    )
    if credibility_count < 6:
        issues.append("事实核查标签少于 6 处")

    if len(re.sub(r"\s+", "", text)) < 2500:
        issues.append("正文过短（少于 2500 字）")
    if "```" in text:
        issues.append("正文残留代码围栏")
    return issues


def _ensure_valid_report(raw: str, topic: str, api_key: str) -> tuple[str, str]:
    issues: list[str] = []
    for attempt in range(MAX_FORMAT_RETRIES + 1):
        text, mermaid = postprocess_report(raw)
        issues = validate_report_structure(text, mermaid)
        if not issues:
            return text, mermaid
        if attempt >= MAX_FORMAT_RETRIES:
            break
        raw = _request_repair(topic, api_key, raw, issues)

    detail = "；".join(issues[:6]) or "未知格式问题"
    raise SkillError(
        f"模型连续 {MAX_FORMAT_RETRIES + 1} 次未通过格式自检，"
        f"已阻止生成不合格 Word。缺失项：{detail}"
    )


# ---------------------------------------------------------------------------
# 黑盒后端：未来群聊系统只需要调用这个函数
# ---------------------------------------------------------------------------


def run_skill(user_input: str, api_key: str) -> SkillResult:
    """
    调用 DeepSeek、清洗正文、提取 Mermaid、生成 Word。

    这个函数不包含任何 Streamlit / UI 代码，也不读写会话状态。
    未来接入多人协作群聊时，可以直接把它放进任务队列或 API 服务：

        result = run_skill(message, room_owner_api_key)
        room.broadcast(result["text"])
        room.attach_file(result["word_bytes"])

    异常：
        SkillError: 输入不合规、API 失败、格式自检失败或 Word 生成失败。
    """
    guard_message = validate_topic(user_input)
    if guard_message:
        raise SkillError(guard_message)

    key = (api_key or "").strip()
    if len(key) < 20:
        raise SkillError("请先在侧边栏填入有效的群主 DeepSeek API Key。")

    topic = user_input.strip()
    raw = _request_report(topic, key)
    text, mermaid = _ensure_valid_report(raw, topic, key)

    try:
        word_bytes = build_word_document(text, mermaid_code=mermaid)
    except Exception as exc:  # noqa: BLE001 - 对外统一转成业务错误
        raise SkillError(f"Word 文档生成失败：{exc}") from exc

    if not word_bytes:
        raise SkillError("Word 文档生成失败：输出为空。")

    return {"text": text, "mermaid_code": mermaid, "word_bytes": word_bytes}


# ---------------------------------------------------------------------------
# python-docx 基础工具
# ---------------------------------------------------------------------------


def _set_run_font(
    run: Run,
    *,
    size: Any = Pt(12),
    bold: bool = False,
    italic: bool = False,
    east_asia: str = BODY_EAST_ASIA,
    ascii_font: str = BODY_ASCII,
    color: RGBColor | None = None,
    underline: bool = False,
) -> None:
    run.font.name = ascii_font
    run.font.size = size
    run.bold = bold
    run.italic = italic
    run.underline = underline
    if color is not None:
        run.font.color.rgb = color

    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)


def _style_font(
    style: Any,
    *,
    east_asia: str,
    ascii_font: str,
    size: int,
    bold: bool = False,
    color: RGBColor | None = None,
) -> None:
    style.font.name = ascii_font
    style.font.size = Pt(size)
    style.font.bold = bold
    if color is not None:
        style.font.color.rgb = color

    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)


def _set_character_spacing(run: Run, twentieths: int) -> None:
    rpr = run._element.get_or_add_rPr()
    existing = rpr.find(qn("w:spacing"))
    if existing is not None:
        rpr.remove(existing)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:val"), str(twentieths))
    rpr.insert_element_before(
        spacing,
        "w:w",
        "w:kern",
        "w:position",
        "w:sz",
        "w:szCs",
        "w:highlight",
        "w:u",
        "w:effect",
    )


def _set_cell_shading(cell: Any, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _set_cell_text(
    cell: Any,
    text: str,
    *,
    bold: bool = False,
    size: Any = Pt(10),
    east_asia: str = BODY_EAST_ASIA,
) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph_format = paragraph.paragraph_format
    paragraph_format.first_line_indent = Pt(0)
    paragraph_format.left_indent = Pt(0)
    paragraph_format.right_indent = Pt(0)
    paragraph_format.space_before = Pt(2)
    paragraph_format.space_after = Pt(2)
    paragraph_format.line_spacing = 1.0
    _add_inline_runs(
        paragraph,
        text.replace("<br>", "\n").replace("<br/>", "\n"),
        size=size,
        bold=bold,
        east_asia=east_asia,
    )


def _set_table_width(table: Any, available_cm: float) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pr = table._tbl.tblPr

    tbl_width = tbl_pr.find(qn("w:tblW"))
    if tbl_width is None:
        tbl_width = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_width)
    tbl_width.set(qn("w:w"), str(int(available_cm / 2.54 * 1440)))
    tbl_width.set(qn("w:type"), "dxa")

    # table.autofit=False 已经写入 tblLayout；这里只补齐缺失的情况，避免重复节点。
    if tbl_pr.find(qn("w:tblLayout")) is None:
        layout = OxmlElement("w:tblLayout")
        layout.set(qn("w:type"), "fixed")
        tbl_pr.append(layout)

    margins = OxmlElement("w:tblCellMar")
    for side, width in (("top", 40), ("left", 80), ("bottom", 40), ("right", 80)):
        element = OxmlElement(f"w:{side}")
        element.set(qn("w:w"), str(width))
        element.set(qn("w:type"), "dxa")
        margins.append(element)

    existing_margins = tbl_pr.find(qn("w:tblCellMar"))
    if existing_margins is not None:
        existing_margins.getparent().replace(existing_margins, margins)
    else:
        look = tbl_pr.find(qn("w:tblLook"))
        if look is not None:
            look.addprevious(margins)
        else:
            tbl_pr.append(margins)


def _plain_inline(text: str) -> str:
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text or "")
    text = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", text)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    text = re.sub(r"(?<!\*)\*(?!\*)", "", text)
    return text.strip()


def _add_hyperlink(
    paragraph: Any,
    url: str,
    label: str,
    *,
    size: Any,
    bold: bool,
    italic: bool,
) -> None:
    part = paragraph.part
    relationship_id = part.relate_to(url, RT.HYPERLINK, is_external=True)

    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    new_run_element = OxmlElement("w:r")
    run = Run(new_run_element, paragraph)
    hyperlink.append(new_run_element)
    paragraph._p.append(hyperlink)

    run.text = label
    _set_run_font(
        run,
        size=size,
        bold=bold,
        italic=italic,
        color=RGBColor(0x05, 0x63, 0xC1),
        underline=True,
    )


def _add_inline_runs(
    paragraph: Any,
    text: str,
    *,
    size: Any = Pt(12),
    bold: bool = False,
    italic: bool = False,
    east_asia: str = BODY_EAST_ASIA,
) -> None:
    """渲染轻量 Markdown 行内语法、可信度标签、链接与换行。"""

    def emit(chunk: str, *, is_bold: bool, is_italic: bool) -> None:
        if not chunk:
            return
        position = 0
        for match in _INLINE_RE.finditer(chunk):
            if match.start() > position:
                run = paragraph.add_run(chunk[position : match.start()])
                _set_run_font(
                    run,
                    size=size,
                    bold=is_bold,
                    italic=is_italic,
                    east_asia=east_asia,
                )

            if match.group("bold"):
                emit(
                    match.group("bold_text"),
                    is_bold=True,
                    is_italic=is_italic,
                )
            elif match.group("italic"):
                emit(
                    match.group("italic_text"),
                    is_bold=is_bold,
                    is_italic=True,
                )
            elif match.group("code"):
                run = paragraph.add_run(match.group("code_text"))
                _set_run_font(
                    run,
                    size=Pt(max(8.5, size.pt - 0.5)),
                    bold=is_bold,
                    italic=is_italic,
                    east_asia="Consolas",
                    ascii_font="Consolas",
                    color=RGBColor(0x7A, 0x1F, 0x55),
                )
            elif match.group("link"):
                _add_hyperlink(
                    paragraph,
                    match.group("link_url"),
                    match.group("link_text"),
                    size=size,
                    bold=is_bold,
                    italic=is_italic,
                )
            elif match.group("cred"):
                level = match.group("cred_level")
                color = {
                    "高": RGBColor(0x0B, 0x6B, 0x3A),
                    "中": RGBColor(0xA8, 0x5B, 0x00),
                    "低": RGBColor(0xB4, 0x1E, 0x1E),
                }.get(level, RGBColor(0x4A, 0x55, 0x68))
                # 用不换行空格锁住“中 ⚠️”，避免标签在 Word 里被拆成两行。
                tag_text = match.group(0).replace(" ", "\u00a0")
                run = paragraph.add_run(tag_text)
                _set_run_font(
                    run,
                    size=size,
                    bold=True,
                    italic=is_italic,
                    east_asia=east_asia,
                    color=color,
                )
            position = match.end()

        if position < len(chunk):
            run = paragraph.add_run(chunk[position:])
            _set_run_font(
                run,
                size=size,
                bold=is_bold,
                italic=is_italic,
                east_asia=east_asia,
            )

    for part in re.split(r"(<br\s*/?>)", text, flags=re.IGNORECASE):
        if re.fullmatch(r"<br\s*/?>", part, flags=re.IGNORECASE):
            paragraph.add_run().add_break()
        else:
            emit(part, is_bold=bold, is_italic=italic)


# ---------------------------------------------------------------------------
# Word 表格渲染引擎（核心）
# ---------------------------------------------------------------------------


def add_markdown_table(doc: Document, md_table_text: str) -> Any | None:
    """
    把 Markdown 管道表转换成真正的 Word 表格。

    - `Table Grid` 边框
    - 表头加粗 + #D9E2F3 浅蓝底
    - 全表 10pt、水平/垂直居中
    - 表头跨页重复、行不跨页拆分
    """
    lines = [line for line in (md_table_text or "").splitlines() if line.strip()]
    if len(lines) < 2:
        return None

    headers = _split_table_row(lines[0])
    if len(lines) > 1 and _is_table_separator(lines[1]):
        data_lines = lines[2:]
    else:
        data_lines = lines[1:]

    rows = [_split_table_row(line) for line in data_lines if line.strip()]
    column_count = max([len(headers), *(len(row) for row in rows)], default=0)
    if column_count == 0:
        return None

    headers += [""] * (column_count - len(headers))
    rows = [row + [""] * (column_count - len(row)) for row in rows]

    table = doc.add_table(rows=1 + len(rows), cols=column_count)
    try:
        table.style = "Table Grid"
    except KeyError:  # pragma: no cover - 极端模板缺样式
        pass

    section = doc.sections[-1]
    available_cm = section.page_width.cm - section.left_margin.cm - section.right_margin.cm
    _set_table_width(table, available_cm)

    def visual_length(value: str) -> int:
        plain = _plain_inline(value)
        return sum(2 if ord(char) > 127 else 1 for char in plain)

    weights: list[int] = []
    for column in range(column_count):
        candidates = [visual_length(headers[column])]
        candidates.extend(visual_length(row[column]) for row in rows)
        weights.append(max(5, min(30, max(candidates))))

    min_width = max(1.4, available_cm / max(column_count, 1) * 0.55)
    widths = [available_cm * weight / sum(weights) for weight in weights]
    widths = [max(min_width, width) for width in widths]

    # 可信度标签和年份/时间是最容易被挤成两行的列，给它们一点最小宽度。
    for column, header in enumerate(headers):
        if "可信度" in header:
            widths[column] = max(widths[column], 3.15)
        elif header.strip() in ("时间", "年份", "阶段"):
            widths[column] = max(widths[column], 2.0)

    scale = available_cm / sum(widths)
    widths = [width * scale for width in widths]

    for column, width in enumerate(widths):
        table.columns[column].width = Cm(width)
        for row in table.rows:
            row.cells[column].width = Cm(width)

    for row_index, row in enumerate(table.rows):
        row_tr_pr = row._tr.get_or_add_trPr()
        cannot_split = OxmlElement("w:cantSplit")
        row_tr_pr.append(cannot_split)
        if row_index == 0:
            repeat_header = OxmlElement("w:tblHeader")
            repeat_header.set(qn("w:val"), "true")
            row_tr_pr.append(repeat_header)

        for column_index, cell in enumerate(row.cells):
            if row_index == 0:
                _set_cell_shading(cell, HEADER_FILL)
                _set_cell_text(
                    cell,
                    headers[column_index],
                    bold=True,
                    east_asia=HEADING_EAST_ASIA,
                )
            else:
                if row_index % 2 == 0:
                    _set_cell_shading(cell, BAND_FILL)
                _set_cell_text(cell, rows[row_index - 1][column_index])

    return table


# ---------------------------------------------------------------------------
# Word 文档排版
# ---------------------------------------------------------------------------


def _configure_document(doc: Document) -> None:
    normal = doc.styles["Normal"]
    _style_font(
        normal,
        east_asia=BODY_EAST_ASIA,
        ascii_font=BODY_ASCII,
        size=12,
        color=RGBColor(0x1B, 0x1F, 0x2A),
    )
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Pt(24)

    heading_specs = (
        ("Heading 1", 16, 16, 8),
        ("Heading 2", 14, 12, 6),
        ("Heading 3", 12, 10, 4),
    )
    for name, size, before, after in heading_specs:
        style = doc.styles[name]
        _style_font(
            style,
            east_asia=HEADING_EAST_ASIA,
            ascii_font=HEADING_EAST_ASIA,
            size=size,
            bold=True,
            color=RGBColor(0x1F, 0x2A, 0x44),
        )
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.25
        style.paragraph_format.keep_with_next = True

    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)
    section.different_first_page_header_footer = True


def _add_field(
    paragraph: Any, instruction: str, placeholder: str | None = "1"
) -> None:
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin.set(qn("w:dirty"), "true")
    begin_run._r.append(begin)

    instruction_run = paragraph.add_run()
    instruction_element = OxmlElement("w:instrText")
    instruction_element.set(qn("xml:space"), "preserve")
    instruction_element.text = f" {instruction} "
    instruction_run._r.append(instruction_element)

    separate_run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(separate)

    if placeholder is not None:
        value_run = paragraph.add_run(placeholder)
        _set_run_font(value_run, size=Pt(9), color=RGBColor(0x66, 0x70, 0x85))

    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def _configure_footers(doc: Document) -> None:
    section = doc.sections[0]

    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    first_paragraph = first_footer.paragraphs[0]
    first_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    first_paragraph.paragraph_format.first_line_indent = Pt(0)
    first_paragraph.paragraph_format.space_before = Pt(0)
    first_paragraph.paragraph_format.space_after = Pt(0)
    run = first_paragraph.add_run("广外 AI 协作空间生成")
    _set_run_font(
        run,
        size=Pt(10.5),
        east_asia=HEADING_EAST_ASIA,
        color=RGBColor(0x7A, 0x84, 0x99),
    )
    _set_character_spacing(run, 20)

    footer = section.footer
    footer.is_linked_to_previous = False
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)

    run = paragraph.add_run("第 ")
    _set_run_font(run, size=Pt(9), color=RGBColor(0x66, 0x70, 0x85))
    _add_field(paragraph, "PAGE")
    run = paragraph.add_run(" 页 / 共 ")
    _set_run_font(run, size=Pt(9), color=RGBColor(0x66, 0x70, 0x85))
    _add_field(paragraph, "NUMPAGES")
    run = paragraph.add_run(" 页")
    _set_run_font(run, size=Pt(9), color=RGBColor(0x66, 0x70, 0x85))


def _add_cover(doc: Document, title: str, subtitle: str, keywords: str) -> None:
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_before = Pt(70)
    spacer.paragraph_format.space_after = Pt(0)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.first_line_indent = Pt(0)
    kicker.paragraph_format.space_after = Pt(26)
    run = kicker.add_run("TOPIC RESEARCH BRIEF  /  深度调研简报")
    _set_run_font(
        run,
        size=Pt(10),
        bold=True,
        east_asia=HEADING_EAST_ASIA,
        color=RGBColor(0x2F, 0x5B, 0xEA),
    )
    _set_character_spacing(run, 40)

    title_paragraph = doc.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.first_line_indent = Pt(0)
    title_paragraph.paragraph_format.space_before = Pt(0)
    title_paragraph.paragraph_format.space_after = Pt(16)
    title_paragraph.paragraph_format.line_spacing = 1.25
    title_run = title_paragraph.add_run(_plain_inline(title))
    _set_run_font(
        title_run,
        size=Pt(24),
        bold=True,
        east_asia=HEADING_EAST_ASIA,
        ascii_font=HEADING_EAST_ASIA,
        color=RGBColor(0x14, 0x1B, 0x33),
    )

    ornament = doc.add_paragraph()
    ornament.alignment = WD_ALIGN_PARAGRAPH.CENTER
    ornament.paragraph_format.first_line_indent = Pt(0)
    ornament.paragraph_format.space_before = Pt(0)
    ornament.paragraph_format.space_after = Pt(20)
    dot_run = ornament.add_run("◆  ◆  ◆")
    _set_run_font(
        dot_run,
        size=Pt(8),
        east_asia=HEADING_EAST_ASIA,
        color=RGBColor(0x2F, 0x5B, 0xEA),
    )
    _set_character_spacing(dot_run, 30)

    subtitle_paragraph = doc.add_paragraph()
    subtitle_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_paragraph.paragraph_format.first_line_indent = Pt(0)
    subtitle_paragraph.paragraph_format.space_after = Pt(28)
    subtitle_run = subtitle_paragraph.add_run(_plain_inline(subtitle))
    _set_run_font(
        subtitle_run,
        size=Pt(16),
        east_asia=HEADING_EAST_ASIA,
        color=RGBColor(0x3D, 0x47, 0x5D),
    )

    keyword_paragraph = doc.add_paragraph()
    keyword_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    keyword_paragraph.paragraph_format.first_line_indent = Pt(0)
    keyword_paragraph.paragraph_format.space_after = Pt(10)
    label_run = keyword_paragraph.add_run("关键词：")
    _set_run_font(
        label_run,
        size=Pt(11),
        bold=True,
        east_asia=HEADING_EAST_ASIA,
        color=RGBColor(0x2F, 0x5B, 0xEA),
    )
    value_run = keyword_paragraph.add_run(keywords or "—")
    _set_run_font(
        value_run,
        size=Pt(11),
        east_asia=BODY_EAST_ASIA,
        color=RGBColor(0x4A, 0x55, 0x68),
    )

    date_paragraph = doc.add_paragraph()
    date_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    date_paragraph.paragraph_format.first_line_indent = Pt(0)
    date_paragraph.paragraph_format.space_before = Pt(8)
    date_run = date_paragraph.add_run(date.today().strftime("%Y年%m月%d日"))
    _set_run_font(
        date_run,
        size=Pt(10),
        east_asia=BODY_EAST_ASIA,
        color=RGBColor(0x8A, 0x93, 0xA6),
    )

    doc.add_page_break()


def _add_toc(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.first_line_indent = Pt(0)
    title.paragraph_format.space_before = Pt(6)
    title.paragraph_format.space_after = Pt(22)
    title_run = title.add_run("目  录")
    _set_run_font(
        title_run,
        size=Pt(18),
        bold=True,
        east_asia=HEADING_EAST_ASIA,
        ascii_font=HEADING_EAST_ASIA,
        color=RGBColor(0x14, 0x1B, 0x33),
    )

    note = doc.add_paragraph()
    note.paragraph_format.first_line_indent = Pt(0)
    note.paragraph_format.space_after = Pt(10)
    note_run = note.add_run("目录将在 Word 中更新域后自动生成。")
    _set_run_font(
        note_run,
        size=Pt(10),
        east_asia=BODY_EAST_ASIA,
        color=RGBColor(0x66, 0x70, 0x85),
    )

    # 空结果域：避免部分阅读器把占位文字渲染成灰色框；
    # Word 打开时会通过 updateFields 自动更新成正式目录。
    toc_paragraph = doc.add_paragraph()
    toc_paragraph.paragraph_format.first_line_indent = Pt(0)
    toc_paragraph.paragraph_format.space_before = Pt(4)
    toc_paragraph.paragraph_format.space_after = Pt(4)
    _add_field(toc_paragraph, 'TOC \\o "1-2" \\h \\z \\u', placeholder=None)

    hint = doc.add_paragraph()
    hint.paragraph_format.first_line_indent = Pt(0)
    hint.paragraph_format.space_before = Pt(14)
    hint_run = hint.add_run(
        "提示：首次打开 Word 时目录可能为空。在目录区域右键，选择“更新域”→“更新整个目录”。"
    )
    _set_run_font(
        hint_run,
        size=Pt(9),
        east_asia=BODY_EAST_ASIA,
        color=RGBColor(0x8A, 0x93, 0xA6),
    )

    doc.add_page_break()


def _add_list_item(doc: Document, text: str, *, ordered: bool) -> None:
    paragraph = doc.add_paragraph()
    paragraph_format = paragraph.paragraph_format
    paragraph_format.left_indent = Pt(24)
    paragraph_format.first_line_indent = Pt(-12)
    paragraph_format.space_before = Pt(0)
    paragraph_format.space_after = Pt(3)
    paragraph_format.line_spacing = 1.5

    if not ordered:
        bullet = paragraph.add_run("• ")
        _set_run_font(
            bullet,
            size=Pt(12),
            bold=True,
            east_asia=BODY_EAST_ASIA,
            color=RGBColor(0x2F, 0x5B, 0xEA),
        )
    _add_inline_runs(paragraph, text)


def _add_blockquote(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph_format = paragraph.paragraph_format
    paragraph_format.left_indent = Pt(18)
    paragraph_format.right_indent = Pt(12)
    paragraph_format.first_line_indent = Pt(0)
    paragraph_format.space_before = Pt(4)
    paragraph_format.space_after = Pt(8)
    paragraph_format.line_spacing = 1.5

    ppr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    left_border = OxmlElement("w:left")
    left_border.set(qn("w:val"), "single")
    left_border.set(qn("w:sz"), "18")
    left_border.set(qn("w:space"), "8")
    left_border.set(qn("w:color"), "2F5BEA")
    borders.append(left_border)
    ppr.insert_element_before(borders, "w:shd", "w:tabs", "w:spacing", "w:ind")

    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), "F4F7FC")
    ppr.insert_element_before(shading, "w:tabs", "w:spacing", "w:ind")

    _add_inline_runs(
        paragraph,
        text,
        size=Pt(12),
        italic=False,
        east_asia=BODY_EAST_ASIA,
    )
    for run in paragraph.runs:
        if run.font.color.rgb is None:
            run.font.color.rgb = RGBColor(0x3D, 0x47, 0x5D)


def _add_heading(doc: Document, text: str, level: int) -> None:
    level = max(1, min(level, 3))
    sizes = {1: Pt(16), 2: Pt(14), 3: Pt(12)}
    paragraph = doc.add_paragraph(style=f"Heading {level}")
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    _add_inline_runs(
        paragraph,
        text,
        size=sizes[level],
        bold=True,
        east_asia=HEADING_EAST_ASIA,
    )
    for run in paragraph.runs:
        run.font.name = HEADING_EAST_ASIA
        run.font.color.rgb = RGBColor(0x1F, 0x2A, 0x44)
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        rfonts.set(qn("w:eastAsia"), HEADING_EAST_ASIA)


def _add_diagram_note(doc: Document) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(8)
    run = paragraph.add_run(
        "（交互式关键词关系图见 Web 成果篮；Word 版以文字说明为准。）"
    )
    _set_run_font(
        run,
        size=Pt(9),
        east_asia=BODY_EAST_ASIA,
        color=RGBColor(0x8A, 0x93, 0xA6),
    )


def _add_body(doc: Document, body: str, *, has_mermaid: bool = False) -> None:
    lines = body.splitlines()
    index = 0
    pending_diagram_note = False

    while index < len(lines):
        raw_line = lines[index].rstrip()
        line = raw_line.strip()
        if not line:
            index += 1
            continue

        if re.fullmatch(r"[-*_—]{3,}", line):
            index += 1
            continue

        if _looks_like_table_row(line) and index + 1 < len(lines) and _is_table_separator(
            lines[index + 1]
        ):
            block: list[str] = []
            while index < len(lines) and _looks_like_table_row(lines[index]):
                block.append(lines[index])
                index += 1
            add_markdown_table(doc, "\n".join(block))
            spacer = doc.add_paragraph()
            spacer.paragraph_format.space_before = Pt(0)
            spacer.paragraph_format.space_after = Pt(4)
            continue

        heading_match = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if heading_match:
            hash_count = len(heading_match.group(1))
            heading_text = re.sub(r"\s+#+\s*$", "", heading_match.group(2)).strip()
            if hash_count == 1:
                index += 1
                continue
            if pending_diagram_note and hash_count == 3:
                _add_diagram_note(doc)
                pending_diagram_note = False
            if hash_count == 2:
                level = 1
            elif hash_count == 3:
                level = 1
            elif hash_count == 4:
                level = 2
            else:
                level = 3
            _add_heading(doc, heading_text, level)
            if has_mermaid and hash_count in (3, 4) and "关键词关系图" in heading_text:
                pending_diagram_note = True
            index += 1
            continue

        if line.startswith(">"):
            quote_lines: list[str] = []
            while index < len(lines) and lines[index].strip().startswith(">"):
                quote_lines.append(
                    re.sub(r"^\s*>\s?", "", lines[index].strip()).strip()
                )
                index += 1
            _add_blockquote(doc, " ".join(quote_lines))
            continue

        bullet_match = re.match(r"^[-*+]\s+(.+)$", line)
        if bullet_match:
            _add_list_item(doc, bullet_match.group(1), ordered=False)
            index += 1
            continue

        ordered_match = re.match(r"^(\d+)[.、)]\s+(.+)$", line)
        if ordered_match:
            _add_list_item(
                doc,
                f"{ordered_match.group(1)}. {ordered_match.group(2)}",
                ordered=True,
            )
            index += 1
            continue

        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.first_line_indent = Pt(24)
        paragraph.paragraph_format.line_spacing = 1.5
        _add_inline_runs(paragraph, line)
        index += 1

    if pending_diagram_note:
        _add_diagram_note(doc)


def _parse_cover(text: str) -> tuple[str, str, str, str]:
    lines = text.splitlines()
    title = ""
    subtitle = ""
    keywords = ""
    body_start = 0

    for index, raw_line in enumerate(lines[:20]):
        line = raw_line.strip()
        if not line:
            body_start = index + 1
            continue
        if not title and line.startswith("# "):
            title = _plain_inline(re.sub(r"\s+#+\s*$", "", line[2:]))
            body_start = index + 1
            continue
        if not subtitle and line.startswith("## "):
            subtitle = _plain_inline(re.sub(r"\s+#+\s*$", "", line[3:]))
            body_start = index + 1
            continue
        keyword_match = _KEYWORDS_RE.match(line)
        if not keywords and keyword_match:
            keywords = _plain_inline(keyword_match.group("value"))
            body_start = index + 1
            continue
        break

    body = "\n".join(lines[body_start:]).lstrip()
    return title, subtitle, keywords, body


def _enable_update_fields(doc: Document) -> None:
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        update = OxmlElement("w:updateFields")
        update.set(qn("w:val"), "true")
        settings.append(update)


def build_word_document(report_text: str, mermaid_code: str = "") -> bytes:
    """把清洗后的 Markdown 报告排版成 Word，返回二进制。"""
    if not isinstance(report_text, str) or not report_text.strip():
        raise ValueError("报告正文为空。")

    title, subtitle, keywords, body = _parse_cover(report_text)
    if not title:
        title = "关键词深度调研报告"
    if not subtitle:
        subtitle = "多源证据、含义变迁与辩论要点"

    doc = Document()
    _configure_document(doc)
    _configure_footers(doc)
    _add_cover(doc, title, subtitle, keywords)
    _add_toc(doc)
    _add_body(doc, body, has_mermaid=bool(mermaid_code))
    _enable_update_fields(doc)

    doc.core_properties.title = title
    doc.core_properties.subject = subtitle
    doc.core_properties.keywords = keywords
    doc.core_properties.author = "广外 AI 协作空间"
    doc.core_properties.comments = (
        "由 skill-searcher_v02 生成；目录需在 Word 中更新域。"
    )
    if mermaid_code:
        doc.core_properties.comments += " 关键词关系图见生成端可视化。"

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# ---------------------------------------------------------------------------
# 成果篮预览辅助
# ---------------------------------------------------------------------------


def _extract_preview(text: str) -> tuple[str, str, str]:
    title_match = re.search(r"(?m)^#\s+(.+)$", text)
    subtitle_match = re.search(r"(?m)^##\s+(.+)$", text)
    title = _plain_inline(title_match.group(1)) if title_match else "关键词深度调研报告"
    subtitle = _plain_inline(subtitle_match.group(1)) if subtitle_match else ""

    lines = text.splitlines()
    summary_lines: list[str] = []
    start_index: int | None = None

    for index, line in enumerate(lines):
        stripped = line.strip()
        if _ABSTRACT_RE.match(stripped):
            start_index = index
            break

    if start_index is not None:
        first = lines[start_index].strip()
        is_quote = first.startswith(">")
        summary_lines.append(first)
        for line in lines[start_index + 1 :]:
            stripped = line.strip()
            if not stripped:
                if summary_lines:
                    break
                continue
            if is_quote and stripped.startswith(">"):
                summary_lines.append(stripped)
                continue
            if not is_quote and not stripped.startswith(("#", "|", ">", "-", "*")):
                summary_lines.append(stripped)
                continue
            break

    summary = "\n".join(summary_lines).strip()
    if not summary:
        for line in lines:
            stripped = line.strip()
            if not stripped or stripped.startswith(("#", ">", "|", "-", "*", "关键词")):
                continue
            summary = f"> **摘要**：{stripped}"
            break

    if len(summary) > 460:
        cutoff = summary[:460]
        sentence_end = cutoff.rfind("。")
        summary = cutoff[: sentence_end + 1] if sentence_end > 200 else cutoff
        summary += "……"
    return title, subtitle, summary


def _source_count(text: str) -> int:
    section = _section_text(text, SECTION_PATTERNS[8][1])
    return len(
        re.findall(r"(?m)^\s*(?:[-*]\s*)?(?:\[\d+\]|\d+[.、])\s*\S", section)
    )


def _credibility_counts(text: str) -> tuple[int, int, int]:
    high = len(re.findall(r"\[可信度\s*[:：]\s*高", text))
    medium = len(re.findall(r"\[可信度\s*[:：]\s*中", text))
    low = len(re.findall(r"\[可信度\s*[:：]\s*低", text))
    return high, medium, low


def _safe_filename(title: str) -> str:
    cleaned = re.sub(r'[\\/:*?"<>|\r\n]+', "_", title).strip(" ._")
    cleaned = re.sub(r"_{2,}", "_", cleaned)[:48] or "关键词深度调研报告"
    return f"{cleaned}_{date.today():%Y%m%d}.docx"


# ---------------------------------------------------------------------------
# 网页层复用：只读预览辅助（不改变 run_skill 行为）
# ---------------------------------------------------------------------------


def report_preview(text: str) -> tuple[str, str, str]:
    """返回 (标题, 副标题, 摘要)，供左侧卡片 / AI 气泡复用。"""
    return _extract_preview(text)


def report_filename(text: str) -> str:
    """根据报告标题生成安全 .docx 文件名（与单机版下载名规则一致）。"""
    title, _subtitle, _summary = _extract_preview(text)
    return _safe_filename(title)


def report_stats(text: str) -> dict[str, int]:
    """报告结构统计：字数 / 表格数 / 来源条数 / 三档可信度数量。"""
    high, medium, low = _credibility_counts(text)
    return {
        "characters": len(re.sub(r"\s+", "", text)),
        "tables": _count_tables(text),
        "sources": _source_count(text),
        "credibility_high": high,
        "credibility_medium": medium,
        "credibility_low": low,
    }
