# -*- coding: utf-8 -*-
"""DeepSeek 总结黑盒（从 skill-summarizer_v01 原样搬入，仅去掉了 Streamlit import）。

``run_summary(chat_history, api_key)`` 的实现、Prompt、Markdown→Word 逻辑
均未改动；本文件只作为 v02 协作空间的复用入口，避免在 UI 层重写模型调用。
"""

from __future__ import annotations

import io
import logging
import re
from datetime import datetime
from typing import Final, TypedDict

import requests
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

logger = logging.getLogger(__name__)


# =============================================================================
# 1. 常量：模型、文件、提示词
# =============================================================================

DEEPSEEK_API_URL: Final[str] = "https://api.deepseek.com/chat/completions"
DEEPSEEK_MODEL: Final[str] = "deepseek-chat"
DEFAULT_FILENAME: Final[str] = "团队复盘纪要.docx"
DOCX_MIME: Final[str] = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)

# DeepSeek 上下文约 64K token。这里给 Prompt 和 4096 token 输出留出余量。
MAX_CHAT_CHARS: Final[int] = 45_000
CHAT_HEAD_CHARS: Final[int] = 10_000
REQUEST_TIMEOUT: Final[tuple[int, int]] = (15, 150)

FONT_NAME: Final[str] = "Microsoft YaHei"
EAST_ASIA_FONT: Final[str] = "微软雅黑"
CODE_FONT: Final[str] = "Consolas"
TEXT_COLOR: Final[RGBColor] = RGBColor(0x1F, 0x29, 0x37)
MUTED_COLOR: Final[RGBColor] = RGBColor(0x64, 0x74, 0x8B)
ACCENT_COLOR: Final[RGBColor] = RGBColor(0x1D, 0x4E, 0xD8)


SYSTEM_PROMPT: Final[str] = """# Role
你是一位资深的项目管家和讨论复盘专家。你的任务是帮助团队理清思路，把杂乱无章的群聊讨论转化为结构化的共识。

# Task
请阅读提供的群聊记录，提取出团队当前的讨论状态。要求极其精准地识别“观点迭代”（即新观点覆盖旧观点的过程）。

# Rules
1. 忽略寒暄、表情包、跑题废话。
2. **状态归类（核心）**：
   - **✅ 肯定（达成共识）**：团队最终决定采用的方案/决定，并提炼出核心原因。
   - **❌ 否定（废弃方案）**：之前讨论过但被推翻、明确否定的方案，并提炼出否定的原因。
   - **⏳ 搁浅（待讨论）**：提出了但目前没有定论、需要进一步查证或留待下次讨论的点。
3. **文件名生成**：根据讨论的核心主题，生成一个概括性的 Word 文件名（例如：1024_迎新活动策划复盘）。

# Output Format
第一行必须是：[FILENAME: 这里填生成的概括性文件名]

然后严格按以下结构输出：
### 🎯 核心讨论主题
[一句话概括本次讨论的核心]
### ✅ 已达成共识（肯定内容及原因）
- **共识1**：[描述]
  - *原因*：[为什么最终定下了这个]
### ❌ 已否决方案（否定内容及原因）
- **废弃方案1**：[描述]
  - *原因*：[为什么被否决了]
### ⏳ 搁浅待议（待讨论的点）
- [尚未有定论的点]

【严厉警告】严禁输出任何寒暄、废话，直接输出规定格式。"""


class SummaryResult(TypedDict):
    """``run_summary`` 的固定返回结构。"""

    text: str
    filename: str
    word_bytes: bytes


class SummaryError(RuntimeError):
    """可直接展示给用户的业务异常。"""


# =============================================================================
# 2. 黑盒业务层：无任何 Streamlit 代码
# =============================================================================

_FILENAME_FIRST_LINE_RE: Final[re.Pattern[str]] = re.compile(
    r"^\[FILENAME:\s*(?P<name>.+?)\s*\]$"
)
_FILENAME_LINE_RE: Final[re.Pattern[str]] = re.compile(
    r"(?m)^[^\n]*\[FILENAME:[^\n]*\][^\n]*(?:\n|$)"
)
_ILLEGAL_FILENAME_RE: Final[re.Pattern[str]] = re.compile(r'[\\/:*?"<>|\x00-\x1f]')
_WINDOWS_RESERVED_NAMES: Final[frozenset[str]] = frozenset(
    {
        "CON",
        "PRN",
        "AUX",
        "NUL",
        *(f"COM{i}" for i in range(1, 10)),
        *(f"LPT{i}" for i in range(1, 10)),
    }
)


def _truncate_chat(chat_history: str, limit: int = MAX_CHAT_CHARS) -> str:
    """清洗空白并做保护性截断，优先保留最近讨论。"""
    cleaned = chat_history.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    cleaned = cleaned.strip()
    if len(cleaned) <= limit:
        return cleaned

    head_budget = min(CHAT_HEAD_CHARS, limit // 4)
    tail_budget = limit - head_budget
    notice = (
        "\n\n【系统提示：原始聊天记录过长，中间部分已截断；"
        "以下保留讨论开头和最近的讨论内容。】\n\n"
    )
    return f"{cleaned[:head_budget]}{notice}{cleaned[-tail_budget:]}"


def _build_user_prompt(chat_history: str) -> str:
    """把聊天记录包进清晰的边界，减少记录误当指令的风险。"""
    return (
        f"请复盘以下群聊记录。\n\n"
        f"当前日期：{datetime.now():%Y-%m-%d}\n\n"
        f"<<<CHAT_START>>>\n{chat_history}\n<<<CHAT_END>>>\n\n"
        f"请严格遵守 System Prompt 的输出格式。"
    )


def _friendly_api_error(response: requests.Response) -> str:
    """把 DeepSeek 的 HTTP 错误转换为新手看得懂的提示。"""
    status = response.status_code
    message = ""
    try:
        payload = response.json()
        error = payload.get("error", {}) if isinstance(payload, dict) else {}
        if isinstance(error, dict):
            message = str(error.get("message", "")).strip()
    except ValueError:
        message = ""

    known_messages = {
        400: "请求格式有误，请稍后重试。",
        401: "API Key 无效或已过期，请检查后重新填写。",
        402: "DeepSeek 账户余额不足，请充值后重试。",
        403: "当前 API Key 没有访问权限。",
        404: "API 地址不存在，请检查接口配置。",
        429: "请求过于频繁或额度受限，请稍后再试。",
        500: "DeepSeek 服务暂时异常，请稍后重试。",
        503: "DeepSeek 服务繁忙，请稍后重试。",
    }
    friendly = known_messages.get(status)
    if friendly and status != 400:
        return friendly
    if message:
        short_message = message[:200]
        return f"{friendly or 'DeepSeek API 请求失败。'}（详情：{short_message}）"
    return friendly or f"DeepSeek API 请求失败（HTTP {status}）。"


def _request_summary(chat_history: str, api_key: str) -> str:
    """请求 DeepSeek 并返回未经清洗的模型文本。"""
    payload = {
        "model": DEEPSEEK_MODEL,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": _build_user_prompt(chat_history)},
        ],
        "temperature": 0.2,
        "max_tokens": 4096,
        "stream": False,
    }
    headers = {
        "Authorization": f"Bearer {api_key}",
        "Content-Type": "application/json",
        "Accept": "application/json",
    }

    try:
        response = requests.post(
            DEEPSEEK_API_URL,
            headers=headers,
            json=payload,
            timeout=REQUEST_TIMEOUT,
        )
    except requests.exceptions.Timeout as exc:
        raise SummaryError("请求超时，请检查网络后重试。") from exc
    except requests.exceptions.RequestException as exc:
        raise SummaryError("无法连接 DeepSeek，请检查网络或代理设置。") from exc

    if response.status_code >= 400:
        raise SummaryError(_friendly_api_error(response))

    try:
        data = response.json()
        content = data["choices"][0]["message"]["content"]
    except (ValueError, KeyError, IndexError, TypeError) as exc:
        raise SummaryError("DeepSeek 返回了无法识别的数据，请稍后重试。") from exc

    if not isinstance(content, str) or not content.strip():
        raise SummaryError("DeepSeek 没有返回有效内容，请重试。")
    return content.strip()


def _strip_outer_code_fence(text: str) -> str:
    """容错处理模型偶尔包上的 ```markdown 代码围栏。"""
    cleaned = text.replace("\ufeff", "").replace("\r\n", "\n").replace("\r", "\n").strip()
    lines = cleaned.split("\n")
    if lines and re.fullmatch(r"```(?:markdown|md)?\s*", lines[0].strip(), re.IGNORECASE):
        lines = lines[1:]
    if lines and lines[-1].strip() == "```":
        lines = lines[:-1]
    return "\n".join(lines).strip()


def _sanitize_filename(raw_name: str) -> str:
    """生成适合 Windows / macOS / Linux 下载的安全 .docx 文件名。"""
    name = raw_name.strip().strip("`'\"“”‘’").strip()
    name = re.sub(r"\.docx?$", "", name, flags=re.IGNORECASE).strip(" .")
    name = _ILLEGAL_FILENAME_RE.sub("_", name)
    name = re.sub(r"\s+", " ", name).strip(" ._")

    if not name:
        name = DEFAULT_FILENAME[:-5]
    if name.upper() in _WINDOWS_RESERVED_NAMES:
        name = f"{name}_复盘纪要"
    if len(name) > 80:
        name = name[:80].rstrip(" ._")

    return f"{name or DEFAULT_FILENAME[:-5]}.docx"


def _parse_model_output(raw_text: str) -> tuple[str, str]:
    """提取文件名并物理删除 FILENAME 行，返回 ``(正文, 文件名)``。"""
    cleaned = _strip_outer_code_fence(raw_text)
    if not cleaned:
        raise SummaryError("DeepSeek 没有返回有效内容，请重试。")

    first_line = cleaned.split("\n", 1)[0].strip().strip("`").strip()
    filename_match = _FILENAME_FIRST_LINE_RE.match(first_line)
    filename = _sanitize_filename(
        filename_match.group("name") if filename_match else DEFAULT_FILENAME
    )

    # 删除所有包含 FILENAME 的行，避免 Markdown 正文被污染。
    body = _FILENAME_LINE_RE.sub("", cleaned).strip()
    body = re.sub(r"\n{3,}", "\n\n", body).strip()

    if not body:
        raise SummaryError("DeepSeek 只返回了文件名，没有返回复盘正文，请重试。")
    return body, filename


# -----------------------------------------------------------------------------
# 2.1 Markdown → Word
# -----------------------------------------------------------------------------

_HEADING_RE: Final[re.Pattern[str]] = re.compile(r"^(?P<hashes>#{1,6})\s+(?P<text>.+)$")
_BULLET_RE: Final[re.Pattern[str]] = re.compile(
    r"^(?P<indent>[ \t]*)[-*+]\s+(?P<text>.+)$"
)
_NUMBER_RE: Final[re.Pattern[str]] = re.compile(
    r"^(?P<indent>[ \t]*)(?P<num>\d+)[.)]\s+(?P<text>.+)$"
)
_BLOCKQUOTE_RE: Final[re.Pattern[str]] = re.compile(r"^>\s?(?P<text>.*)$")
_INLINE_RE: Final[re.Pattern[str]] = re.compile(
    r"\*\*\*(?P<bi>(?!\s)[^*\n]*?(?<!\s))\*\*\*"
    r"|\*\*(?P<b>(?!\s)[^*\n]*?(?<!\s))\*\*"
    r"|\*(?P<i>(?!\s)[^*\n]*?(?<!\s))\*"
    r"|`(?P<c>[^`\n]+?)`"
)


def _set_east_asia_font(rpr_element: object, font_name: str) -> None:
    """给 run / style 的 rPr 节点补上东亚字体声明。"""
    rpr = rpr_element  # 保持调用处类型简单；python-docx 的 rPr 是 lxml 节点。
    rfonts = rpr.find(qn("w:rFonts"))  # type: ignore[attr-defined]
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)  # type: ignore[attr-defined]
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia"):
        rfonts.set(qn(attribute), font_name)


def _style_run(
    run: object,
    *,
    bold: bool = False,
    italic: bool = False,
    code: bool = False,
) -> None:
    """统一设置单个文本片段的字体、字号和强调样式。"""
    run.font.name = CODE_FONT if code else FONT_NAME  # type: ignore[attr-defined]
    run.font.size = Pt(9.5) if code else Pt(10.5)  # type: ignore[attr-defined]
    run.bold = bold  # type: ignore[attr-defined]
    run.italic = italic  # type: ignore[attr-defined]
    _set_east_asia_font(run._element.get_or_add_rPr(), CODE_FONT if code else EAST_ASIA_FONT)  # type: ignore[attr-defined]


def _add_inline_runs(paragraph: object, text: str, *, base_italic: bool = False) -> None:
    """解析 ``**粗体**``、``*斜体*`` 与行内代码，并写入 Word 段落。"""
    cursor = 0
    for match in _INLINE_RE.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor : match.start()])  # type: ignore[attr-defined]
            _style_run(run, italic=base_italic)
        if match.group("bi") is not None:
            _style_run(
                paragraph.add_run(match.group("bi")),  # type: ignore[attr-defined]
                bold=True,
                italic=True,
            )
        elif match.group("b") is not None:
            _style_run(
                paragraph.add_run(match.group("b")),  # type: ignore[attr-defined]
                bold=True,
                italic=base_italic,
            )
        elif match.group("i") is not None:
            _style_run(
                paragraph.add_run(match.group("i")),  # type: ignore[attr-defined]
                italic=True,
            )
        elif match.group("c") is not None:
            _style_run(
                paragraph.add_run(match.group("c")),  # type: ignore[attr-defined]
                code=True,
                italic=base_italic,
            )
        cursor = match.end()

    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])  # type: ignore[attr-defined]
        _style_run(run, italic=base_italic)


def _add_paragraph_shading(paragraph: object, fill: str) -> None:
    """给段落加底色（用于代码块）。"""
    p_pr = paragraph._p.get_or_add_pPr()  # type: ignore[attr-defined]
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def _add_code_block(document: object, lines: list[str]) -> None:
    """把代码块按原样写入一个带浅灰底色的段落。"""
    if not lines:
        return
    paragraph = document.add_paragraph()  # type: ignore[attr-defined]
    paragraph.paragraph_format.left_indent = Cm(0.5)
    paragraph.paragraph_format.right_indent = Cm(0.5)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(6)
    _add_paragraph_shading(paragraph, "F1F5F9")

    for index, code_line in enumerate(lines):
        if index:
            paragraph.add_run().add_break()
        run = paragraph.add_run(code_line)
        _style_run(run, code=True)


def _bullet_style(document: object, level: int) -> str:
    """返回当前 Word 模板中可用的列表样式名。"""
    style_name = "List Bullet" if level <= 0 else f"List Bullet {min(level + 1, 3)}"
    try:
        document.styles[style_name]  # type: ignore[attr-defined]
    except KeyError:
        return "List Bullet"
    return style_name


def _configure_document(document: Document, title: str) -> None:
    """设置页面、字体、标题样式与页脚。"""
    section = document.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    normal = document.styles["Normal"]
    normal.font.name = FONT_NAME
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = TEXT_COLOR
    _set_east_asia_font(normal.element.get_or_add_rPr(), EAST_ASIA_FONT)
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(4)

    heading_specs = (
        ("Heading 1", 17, 16),
        ("Heading 2", 14, 12),
        ("Heading 3", 12.5, 10),
    )
    for style_name, size, space_before in heading_specs:
        style = document.styles[style_name]
        style.font.name = FONT_NAME
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = ACCENT_COLOR
        _set_east_asia_font(style.element.get_or_add_rPr(), EAST_ASIA_FONT)
        style.paragraph_format.space_before = Pt(space_before)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True

    document.core_properties.title = title
    document.core_properties.author = "群聊通用复盘总结器"
    document.core_properties.comments = "由 DeepSeek 辅助生成的团队讨论复盘纪要。"

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_run = footer.add_run(
        f"{title} · {datetime.now():%Y-%m-%d} · 由群聊通用复盘总结器生成"
    )
    footer_run.font.name = FONT_NAME
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = MUTED_COLOR
    _set_east_asia_font(footer_run._element.get_or_add_rPr(), EAST_ASIA_FONT)


def _markdown_to_docx(markdown_text: str, filename: str) -> bytes:
    """把纯 Markdown 文本转换为可下载的 Word 文档字节流。"""
    document = Document()
    title = re.sub(r"\.docx$", "", filename, flags=re.IGNORECASE)
    _configure_document(document, title)

    lines = markdown_text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    heading_levels = [
        len(match.group("hashes"))
        for line in lines
        if (match := _HEADING_RE.match(line.strip()))
    ]
    base_heading_level = min(heading_levels) if heading_levels else 1

    in_code_block = False
    code_lines: list[str] = []

    for raw_line in lines:
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code_block:
                _add_code_block(document, code_lines)
                code_lines = []
                in_code_block = False
            else:
                in_code_block = True
            continue
        if in_code_block:
            code_lines.append(line)
            continue

        if not stripped:
            continue
        if re.fullmatch(r"[-*_]{3,}", stripped):
            continue

        heading_match = _HEADING_RE.match(stripped)
        if heading_match:
            relative_level = len(heading_match.group("hashes")) - base_heading_level + 1
            level = max(1, min(relative_level, 3))
            paragraph = document.add_paragraph(style=f"Heading {level}")
            _add_inline_runs(paragraph, heading_match.group("text").strip())
            continue

        if _BLOCKQUOTE_RE.match(stripped):
            quote_text = _BLOCKQUOTE_RE.match(stripped).group("text")  # type: ignore[union-attr]
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.6)
            _add_inline_runs(paragraph, quote_text, base_italic=True)
            continue

        bullet_match = _BULLET_RE.match(line)
        if bullet_match:
            indent = bullet_match.group("indent").replace("\t", "  ")
            level = min(len(indent) // 2, 2)
            paragraph = document.add_paragraph(style=_bullet_style(document, level))
            paragraph.paragraph_format.space_after = Pt(2)
            _add_inline_runs(paragraph, bullet_match.group("text").strip())
            continue

        number_match = _NUMBER_RE.match(line)
        if number_match:
            indent = number_match.group("indent").replace("\t", "  ")
            level = min(len(indent) // 2, 2)
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.75 + level * 0.6)
            paragraph.paragraph_format.first_line_indent = Cm(-0.4)
            paragraph.paragraph_format.space_after = Pt(2)
            number_run = paragraph.add_run(f"{number_match.group('num')}. ")
            _style_run(number_run, bold=True)
            _add_inline_runs(paragraph, number_match.group("text").strip())
            continue

        paragraph = document.add_paragraph()
        _add_inline_runs(paragraph, stripped)

    if in_code_block:
        _add_code_block(document, code_lines)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def markdown_to_docx(markdown_text: str, filename: str = "报告.docx") -> bytes:
    """把任意 Markdown 报告渲染为 Word 文档，供成果篮的「诊断与报告」文件夹使用。"""
    text = (markdown_text or "").strip()
    if not text:
        raise ValueError("报告内容为空")
    safe_name = _sanitize_filename(filename)
    return _markdown_to_docx(text, safe_name)


def run_summary(chat_history: str, api_key: str) -> dict:
    """把群聊记录变成结构化复盘文本和 Word 文件。

    Args:
        chat_history: 原始群聊记录（可包含房间名等上下文）。
        api_key: DeepSeek API Key，仅用于本次请求。

    Returns:
        固定包含 ``text``、``filename``、``word_bytes`` 三个字段的字典。

    Raises:
        SummaryError: 输入为空、API 失败或模型输出无效时抛出。
    """
    chat = chat_history.strip()
    key = api_key.strip()
    if not chat:
        raise SummaryError("请输入聊天记录")
    if not key:
        raise SummaryError("请填写 DeepSeek API Key")

    prepared_chat = _truncate_chat(chat)
    raw_text = _request_summary(prepared_chat, key)
    text, filename = _parse_model_output(raw_text)
    word_bytes = _markdown_to_docx(text, filename)

    return {
        "text": text,
        "filename": filename,
        "word_bytes": word_bytes,
    }
